"""Additive UK item-level CPI decomposition.

This module deliberately sits beside the existing five-bucket decomposition,
the same way danish_item_decomposition.py does for Denmark. It selects a
non-overlapping COICOP frontier from ONS's detailed MM23 CPI INDEX series,
using ONS's separately published Annex A "Tables W1 to W3" weight workbook
(committed as a reference file per year, the same convention as Denmark's
weight workbook -- ONS does not serve it through a stable API). Frontier
selection cascades a residual at every level of the tree, using a node's
own rate for whatever its resolved children don't cover, rather than only
reconciling once at the very top -- see select_uk_frontier.
"""

import re
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

from inflation_engine.utils.io import atomic_write_parquet
from inflation_engine.utils.paths import (
    ONS_ANNEX_A_WEIGHTS_DIR,
    ONS_UK_CPI_PATH,
    UK_ITEM_DECOMPOSITION_PATH,
)


OUTPUT_COLUMNS = [
    "date",
    "country",
    "country_code",
    "source_component_code",
    "coicop_code",
    "description",
    "hierarchy_level",
    "series_role",
    "weight_pct",
    "weight_reference_date",
    "weight_basis_year",
    "index_value",
    "inflation_yoy",
    "estimated_contribution_yoy_pp",
    "headline_yoy",
    "reconciliation_row",
]

# The workbook's own CDID namespace for weights is unrelated to the CDID
# namespace fetch_ons.py uses for index series, so weights are matched by the
# COICOP code text embedded in each row's description, not by CDID.
_ITEM_LABEL_PATTERN = re.compile(
    r"^\s*([0-9]+(?:\.[0-9]+)*(?:/[0-9]+)?[A-Za-z]?)\s+(.+?)\s*$"
)
# ONS's own weight basis years are published Feb-Dec; January uses the prior
# year's weights and is deliberately left unmatched, mirroring the existing
# rationale in fetch_ons_weights.py / docs/decomposition.md.
_MIN_WEIGHT_YEAR = 2023
_MAX_WEIGHT_YEAR = 2026

# fetch_ons.py's seven curated special-aggregate series use their native ONS
# CDID as source_series_id (not a COICOP code, unlike the bulk-discovered
# items). Of those seven, only actual rents for housing also corresponds to a
# single COICOP frontier node (04.1) and needs bridging to join weights.
_CURATED_COICOP_OVERRIDES = {"D7CE": "04.1"}


def _feb_dec_column(header_row: tuple, year: int) -> int:
    pattern = re.compile(rf"^{year}\d?\s*\n\s*Feb-Dec\s*$")
    for index, value in enumerate(header_row):
        if isinstance(value, str) and pattern.match(value.strip()):
            return index
    raise ValueError(f"No Feb-Dec ONS weight column found for {year}")


def load_uk_weight_basis(
    year: int,
    weights_dir: str | Path = ONS_ANNEX_A_WEIGHTS_DIR,
) -> pd.DataFrame:
    """Parse one year's W1-CPI sheet from ONS's Annex A weight workbook."""
    path = Path(weights_dir) / f"annexa_w1w3_{year}.xlsx"
    if not path.is_file():
        raise FileNotFoundError(f"No ONS Annex A weight workbook for {year}: {path}")
    workbook = openpyxl.load_workbook(path, data_only=True)
    sheet = workbook["W1-CPI"]
    rows = list(sheet.iter_rows(min_row=1, max_row=sheet.max_row, values_only=True))
    column = _feb_dec_column(rows[4], year)

    records = []
    for row in rows[5:]:
        cdid, description = row[1], row[2]
        if not isinstance(cdid, str) or not isinstance(description, str):
            continue
        if not re.fullmatch(r"[A-Z0-9]{4,6}", cdid.strip()):
            continue
        weight_pct = pd.to_numeric(row[column], errors="coerce")
        if pd.isna(weight_pct):
            continue
        stripped = description.strip()
        if stripped.upper() == "CPI (OVERALL INDEX)":
            code, label = "00", "All items"
        else:
            match = _ITEM_LABEL_PATTERN.match(description)
            if not match:
                continue
            code, label = match.group(1), match.group(2)
        records.append(
            {
                "cdid": cdid.strip(),
                "coicop_code": code,
                "description": label,
                "weight_pct": float(weight_pct) / 10.0,
            }
        )
    weights = pd.DataFrame(records)
    if weights.empty:
        raise ValueError(f"Parsed no ONS weight rows for {year}")
    # The sheet re-lists each division header a second time (e.g. both "10"
    # and "10.0" carry CDID CJUU) right before its detailed breakdown; keep
    # only the first occurrence of each CDID to avoid double-counting it.
    weights = weights.drop_duplicates("cdid", keep="first").reset_index(drop=True)
    weights["weight_reference_date"] = pd.Timestamp(year, 2, 1)
    weights["weight_basis_year"] = year
    return weights


def _implied_prefixes(code: str) -> list[str]:
    """Expand a compressed code like '08.2/3' into ['08.2', '08.3'], or
    '01.1.9.1/2' into ['01.1.9.1', '01.1.9.2']. A plain code returns itself.
    ONS uses '/' to combine sibling groups into one published row; some of
    those siblings are *also* published individually at finer detail
    elsewhere in the same sheet, which this makes it possible to detect."""
    if "/" not in code:
        return [code]
    head, _, alt = code.rpartition(".")
    base, _, rest = alt.partition("/")
    prefixes = [f"{head}.{base}" if head else base]
    prefixes.extend(f"{head}.{extra}" if head else extra for extra in rest.split("/"))
    return prefixes


_UK_WEIGHT_TOLERANCE = 1e-9


def _direct_parent(code: str, codes: set[str]) -> str | None:
    """The nearest ancestor of ``code`` among ``codes``, accounting for
    compressed range codes (e.g. "08.2/3" is a valid parent of "08.2.0.1"
    via its "08.2" branch). Ties (multiple ancestors matching) are broken
    by the longest matched prefix, i.e. the closest ancestor."""
    best: str | None = None
    best_len = -1
    for other in codes:
        if other == code:
            continue
        for prefix in _implied_prefixes(other):
            if code.startswith(prefix + ".") and len(prefix) > best_len:
                best, best_len = other, len(prefix)
    return best


def select_uk_frontier(
    weights: pd.DataFrame,
    available_component_codes: set[str],
) -> pd.DataFrame:
    """Select the deepest available weight nodes with a matched index series.

    ONS frequently publishes a weight and index for a parent node whose
    listed children only partially cover it -- e.g. "10 Education" (weight
    3.42) has only one published child, "10.4 Tertiary education" (2.35);
    the ~1.07pp gap has no separate line item, but Education's own index
    (D7C5, already fetched) can stand in for that remainder. This cascades
    a residual at every level, using a node's own rate for whatever its
    resolved children don't cover, rather than dropping straight to the
    single top-level reconciliation row.
    """
    matched = weights.loc[weights["coicop_code"].ne("00")].copy()
    codes = set(matched["coicop_code"])
    weight_by_code = matched.set_index("coicop_code")["weight_pct"]
    row_by_code = matched.set_index("coicop_code")

    parent_of = {code: _direct_parent(code, codes) for code in codes}
    children_of: dict[str, list[str]] = {code: [] for code in codes}
    for code, parent in parent_of.items():
        if parent is not None:
            children_of[parent].append(code)

    def resolve(code: str) -> list[tuple[str, float]]:
        parts: list[tuple[str, float]] = []
        covered = 0.0
        for child in children_of[code]:
            child_parts = resolve(child)
            parts.extend(child_parts)
            covered += sum(weight for _, weight in child_parts)
        remainder = weight_by_code[code] - covered
        if remainder > _UK_WEIGHT_TOLERANCE and code in available_component_codes:
            parts.append((code, remainder))
        return parts

    top_level = [code for code in codes if parent_of[code] is None]
    parts: list[tuple[str, float]] = []
    for top in top_level:
        parts.extend(resolve(top))
    if not parts:
        raise ValueError("No UK weight rows survived frontier selection")

    frontier = pd.DataFrame(parts, columns=["coicop_code", "weight_pct"])
    full_weight = frontier["coicop_code"].map(weight_by_code)
    is_partial = (full_weight - frontier["weight_pct"]) > _UK_WEIGHT_TOLERANCE
    frontier["cdid"] = frontier["coicop_code"].map(row_by_code["cdid"])
    frontier["description"] = frontier["coicop_code"].map(row_by_code["description"])
    frontier.loc[is_partial, "description"] += " (unitemized remainder only)"
    frontier["weight_reference_date"] = frontier["coicop_code"].map(
        row_by_code["weight_reference_date"]
    )
    frontier["weight_basis_year"] = frontier["coicop_code"].map(
        row_by_code["weight_basis_year"]
    )
    frontier["hierarchy_level"] = frontier["coicop_code"].str.count(r"\.") + 1
    return frontier


def _validate_full_basket_coverage(frontier: pd.DataFrame, year: int) -> None:
    """Sanity-check that a frontier built from the full annual weight basis
    covers essentially the whole basket. select_uk_frontier's cascading
    residuals recover most gaps where ONS's published detail doesn't fully
    itemize a parent's own weight (e.g. Education), so only genuinely
    unmatched subtrees (no index anywhere from leaf to top) should be
    missing here -- a much larger gap would mean a broken parse, not a
    tolerable source-data gap."""
    total = float(frontier["weight_pct"].sum())
    if not (85.0 <= total <= 100.5):
        raise ValueError(
            f"UK component frontier for {year} weight basis sums to "
            f"{total:.2f}%, outside the expected 85-100.5% range; refusing "
            "to build on a broken parse"
        )


def build_uk_item_decomposition(
    raw_cpi: pd.DataFrame,
    weights_dir: str | Path = ONS_ANNEX_A_WEIGHTS_DIR,
) -> pd.DataFrame:
    """Build UK item contributions while leaving the canonical buckets untouched."""
    required = {
        "date", "country", "country_code", "series_role", "component",
        "source_series_id", "value",
    }
    missing = required.difference(raw_cpi.columns)
    if missing:
        raise ValueError(f"UK item panel is missing columns: {sorted(missing)}")

    source = raw_cpi.loc[
        (raw_cpi["country_code"] == "GB") & (raw_cpi["series_role"] == "component")
    ].copy()
    source["date"] = pd.to_datetime(source["date"])
    source["source_component_code"] = (
        source["source_series_id"]
        .astype(str)
        .map(lambda code: _CURATED_COICOP_OVERRIDES.get(code, code))
    )
    source["index_value"] = pd.to_numeric(source["value"], errors="coerce")
    source = source.sort_values(["source_component_code", "date"])
    source["inflation_yoy"] = source.groupby("source_component_code")[
        "index_value"
    ].pct_change(12) * 100.0

    headline = raw_cpi.loc[
        (raw_cpi["country_code"] == "GB") & (raw_cpi["source_series_id"] == "D7BT")
    ].copy()
    headline["date"] = pd.to_datetime(headline["date"])
    headline["index_value"] = pd.to_numeric(headline["value"], errors="coerce")
    headline = headline.sort_values("date")
    headline["headline_yoy"] = headline["index_value"].pct_change(12) * 100.0
    headline = headline[["date", "headline_yoy"]]

    available_codes = set(source["source_component_code"].unique())
    selected_frames = []
    for year in range(_MIN_WEIGHT_YEAR, _MAX_WEIGHT_YEAR + 1):
        weights = load_uk_weight_basis(year, weights_dir)
        frontier = select_uk_frontier(weights, available_codes)
        _validate_full_basket_coverage(frontier, year)
        year_source = source.loc[
            (source["date"].dt.year == year) & (source["date"].dt.month >= 2)
        ]
        merged = year_source.merge(
            frontier,
            left_on="source_component_code",
            right_on="coicop_code",
            how="inner",
            validate="many_to_one",
        )
        selected_frames.append(merged)
    if not selected_frames:
        raise ValueError("No UK weight-basis years matched detailed CPI components")
    selected = pd.concat(selected_frames, ignore_index=True)
    selected["estimated_contribution_yoy_pp"] = (
        selected["weight_pct"] * selected["inflation_yoy"] / 100.0
    )
    selected = selected.merge(headline, on="date", how="left", validate="many_to_one")
    selected["reconciliation_row"] = False

    sums = selected.groupby("date")["estimated_contribution_yoy_pp"].sum(min_count=1)
    target = headline.set_index("date")["headline_yoy"]
    reconciliation = (target - sums).rename("estimated_contribution_yoy_pp")
    reconciliation = reconciliation.dropna().reset_index()
    reconciliation["country"] = "United Kingdom"
    reconciliation["country_code"] = "GB"
    reconciliation["source_component_code"] = "__reconciliation__"
    reconciliation["coicop_code"] = "__reconciliation__"
    reconciliation["description"] = "Unallocated reconciliation to headline CPI"
    reconciliation["hierarchy_level"] = 0
    reconciliation["series_role"] = "reconciliation"
    reconciliation["weight_pct"] = 0.0
    reconciliation["weight_reference_date"] = reconciliation["date"]
    reconciliation["weight_basis_year"] = reconciliation["date"].dt.year
    reconciliation["index_value"] = np.nan
    reconciliation["inflation_yoy"] = np.nan
    reconciliation["headline_yoy"] = reconciliation["date"].map(target)
    reconciliation["reconciliation_row"] = True

    output = pd.concat([selected, reconciliation], ignore_index=True, sort=False)
    output["country"] = output["country"].fillna("United Kingdom")
    output["country_code"] = "GB"
    return output.loc[:, OUTPUT_COLUMNS].sort_values(
        ["date", "reconciliation_row", "coicop_code"], kind="stable"
    ).reset_index(drop=True)


def save_uk_item_decomposition(
    frame: pd.DataFrame,
    output_path: str | Path = UK_ITEM_DECOMPOSITION_PATH,
) -> Path:
    return atomic_write_parquet(frame, output_path, columns=OUTPUT_COLUMNS)


def run(
    input_path: str | Path = ONS_UK_CPI_PATH,
    weights_dir: str | Path = ONS_ANNEX_A_WEIGHTS_DIR,
    output_path: str | Path = UK_ITEM_DECOMPOSITION_PATH,
) -> Path:
    frame = pd.read_parquet(input_path)
    return save_uk_item_decomposition(
        build_uk_item_decomposition(frame, weights_dir), output_path
    )


if __name__ == "__main__":
    destination = run()
    print(f"Saved UK item decomposition to {destination}")
